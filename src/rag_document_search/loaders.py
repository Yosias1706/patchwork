"""Safe local document loading with optional PDF and DOCX support."""

from __future__ import annotations

import json
from collections.abc import Iterable
from hashlib import sha256
from html.parser import HTMLParser
from pathlib import Path

from .models import SourceDocument

SUPPORTED_SUFFIXES = {".txt", ".md", ".rst", ".html", ".htm", ".json", ".pdf", ".docx"}


class _HTMLTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)

    @property
    def text(self) -> str:
        return " ".join(part.strip() for part in self.parts if part.strip())


def load_path(path: Path, *, source: str | None = None) -> SourceDocument:
    """Read one supported document, keeping its origin and stable content metadata."""
    path = path.expanduser().resolve()
    if not path.is_file():
        raise FileNotFoundError(f"Not a file: {path}")
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED_SUFFIXES:
        raise ValueError(f"Unsupported file type: {suffix or '(no extension)'}")

    if suffix == ".pdf":
        text = _read_pdf(path)
    elif suffix == ".docx":
        text = _read_docx(path)
    else:
        raw = path.read_text(encoding="utf-8", errors="replace")
        text = _convert_text(raw, suffix)

    if not text.strip():
        raise ValueError(f"No readable text found in {path.name}")
    return SourceDocument(
        source=source or str(path),
        text=text.strip(),
        content_type=_content_type(suffix),
        metadata={
            "filename": path.name,
            "suffix": suffix,
            "content_sha256": sha256(text.encode("utf-8")).hexdigest(),
            "size_bytes": path.stat().st_size,
        },
    )


def iter_supported_paths(directory: Path) -> Iterable[Path]:
    directory = directory.expanduser().resolve()
    if not directory.is_dir():
        raise NotADirectoryError(f"Not a directory: {directory}")
    for path in sorted(directory.rglob("*")):
        if path.suffix.lower() in SUPPORTED_SUFFIXES:
            yield path


def _convert_text(raw: str, suffix: str) -> str:
    if suffix in {".html", ".htm"}:
        parser = _HTMLTextExtractor()
        parser.feed(raw)
        return parser.text
    if suffix == ".json":
        return json.dumps(json.loads(raw), ensure_ascii=False, indent=2, sort_keys=True)
    return raw


def _read_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as error:  # pragma: no cover - depends on optional dependency
        raise RuntimeError("PDF support needs `pip install -e .[documents]`") from error
    return "\n\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def _read_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as error:  # pragma: no cover - depends on optional dependency
        raise RuntimeError("DOCX support needs `pip install -e .[documents]`") from error
    return "\n".join(paragraph.text for paragraph in Document(path).paragraphs)


def _content_type(suffix: str) -> str:
    return {
        ".txt": "text/plain",
        ".md": "text/markdown",
        ".rst": "text/x-rst",
        ".html": "text/html",
        ".htm": "text/html",
        ".json": "application/json",
        ".pdf": "application/pdf",
        ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }[suffix]
